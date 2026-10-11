from __future__ import annotations

import importlib.util
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document
from lxml.html import fragment_fromstring
from PIL import Image

from caption import answer_export, export_ui, markdown_html, math_render
from caption.chat_tab import ChatPanel
from caption.conversations import ChatMessage
from caption.gui import App
from caption.vocab import NewWord
from caption.vocab_tab import VocabPanel, VocabTranslation


SAMPLE = r"""# Giới hạn và công thức

Chữ **đậm**, *nghiêng* và `code`. Xét \\(x\_1\\) và \\(L\\).

\\[\boxed{\lim\_{x\to a}f(x)=\frac{1}{3}}\\]

| Hàm số | Kết quả |
| --- | --- |
| $x^2$ | $\frac{b^3}{3}$ |
| Hằng số | Tiếng Việt đầy đủ |

3. Mục đầu
   - Mục con
4. Mục tiếp

> Nội dung trích dẫn.

```python
def example(x):
    return x < 3 and x > 0
```

[Tài liệu](https://example.com/docs)
"""


class AnswerExportTests(unittest.TestCase):
    def test_print_math_has_high_resolution_without_changing_screen_dpi(self):
        screen_dpi = math_render.dpi
        root = fragment_fromstring(answer_export.prepare_html(r"\\[x\_1^2\\]"))
        image = root.xpath(".//img")[0]
        data, width, _ = answer_export._image(image, 504)
        with Image.open(io.BytesIO(data)) as png:
            self.assertAlmostEqual(png.width / width, 300 / 72)
        self.assertEqual(math_render.dpi, screen_dpi)

    def test_docx_preserves_text_tables_styles_links_and_formulas(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "answer.docx"
            answer_export.export_html(answer_export.prepare_html(SAMPLE), path)
            doc = Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("Giới hạn và công thức", text)
            self.assertIn("3. Mục đầu", text)
            self.assertIn("• Mục con", text)
            self.assertIn("return x < 3 and x > 0", text)
            self.assertEqual(len(doc.tables), 1)
            self.assertEqual(doc.tables[0].cell(2, 1).text, "Tiếng Việt đầy đủ")
            self.assertEqual(len(doc.inline_shapes), 5)
            self.assertTrue(any(r.bold for p in doc.paragraphs for r in p.runs))
            self.assertTrue(any(r.italic for p in doc.paragraphs for r in p.runs))
            self.assertTrue(any(r.target_ref == "https://example.com/docs" for r in doc.part.rels.values()))
            self.assertIn('w:tblHeader', doc.tables[0]._tbl.xml)

    def test_pdf_is_valid_and_contains_vietnamese(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "answer.pdf"
            answer_export.export_html(answer_export.prepare_html(SAMPLE), path)
            self.assertTrue(path.read_bytes().startswith(b"%PDF-"))
            if importlib.util.find_spec("pypdf"):
                from pypdf import PdfReader
                reader = PdfReader(path)
                text = "\n".join(p.extract_text() for p in reader.pages)
                self.assertIn("Giới hạn và công thức", text)
                self.assertIn("Tiếng Việt đầy đủ", text)
                self.assertIn("Mục con", text)
                self.assertIn("return x < 3 and x > 0", text)

    def test_long_table_and_code_can_paginate_without_losing_last_row(self):
        content = "| Index | Text |\n|---|---|\n" + "\n".join(f"| {n} | Nội dung {n} |" for n in range(100))
        content += "\n\n```\n" + "long_identifier_" * 40 + "\nEND_OF_CODE\n```"
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "long.pdf"
            answer_export.export_html(answer_export.prepare_html(content), path)
            if importlib.util.find_spec("pypdf"):
                from pypdf import PdfReader
                reader = PdfReader(path)
                self.assertGreater(len(reader.pages), 1)
                text = "\n".join(p.extract_text() for p in reader.pages)
                self.assertIn("Nội dung 99", text)
                self.assertIn("END_OF_CODE", text)

    def test_export_failure_preserves_existing_file_and_cleans_temporary_file(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "answer.docx"
            path.write_bytes(b"original")
            with patch.object(answer_export, "_docx", side_effect=RuntimeError("conversion failed")):
                with self.assertRaisesRegex(RuntimeError, "conversion failed"):
                    answer_export.export_html("<p>answer</p>", path)
            self.assertEqual(path.read_bytes(), b"original")
            self.assertEqual(list(Path(folder).iterdir()), [path])

    def test_remote_and_local_images_are_not_read(self):
        content = '<p>A<img src="https://example.com/pic.png" alt="Remote picture"><img src="file:///secret.png" alt="Local picture"></p>'
        with tempfile.TemporaryDirectory() as folder, patch("urllib.request.urlopen") as fetch:
            for extension in (".docx", ".pdf"):
                answer_export.export_html(content, Path(folder) / ("answer" + extension))
            doc = Document(Path(folder) / "answer.docx")
        fetch.assert_not_called()
        self.assertIn("Remote pictureLocal picture", doc.paragraphs[0].text)

    def test_invalid_extension_and_formula_failure_are_reported(self):
        with self.assertRaises(answer_export.ExportError):
            answer_export.export_html("<p>answer</p>", Path("answer.txt"))
        with patch.object(markdown_html, "to_html", return_value="<code class='tex'>bad</code>"):
            with self.assertRaises(answer_export.ExportError):
                answer_export.prepare_html("bad formula")


class AnswerExportUITests(unittest.TestCase):
    def test_chat_exports_exact_raw_answer_and_defers_link_callback(self):
        panel = object.__new__(ChatPanel)
        panel._messages = []
        panel._get_conversation_id = Mock(return_value="conversation")
        panel._store = Mock()
        panel._store.add_message.return_value = 123
        panel._render = Mock()
        panel.after_idle = Mock()
        panel._add_message("bot", "ChatGPT", "**Original answer**", markdown=True)
        key = panel._messages[-1][3]
        panel._on_history_link("rtc-answer-export:" + key)
        with patch.object(export_ui, "choose_export") as choose:
            choose.assert_not_called()
            panel.after_idle.call_args.args[0]()
        choose.assert_called_once_with(panel, "**Original answer**", f"chatgpt-{key}")

    def test_saved_chat_answers_receive_export_links_after_reopening(self):
        panel = object.__new__(ChatPanel)
        panel._clear_image = panel.refresh_state = panel._render = Mock()
        panel._messages, panel._requests = [], {}
        panel._get_conversation_id = Mock(return_value="conversation")
        panel._store = Mock()
        panel._store.messages.return_value = [
            ChatMessage(1, "conversation", "bot", "chat", "ChatGPT", "first", None, "2026-10-09"),
            ChatMessage(2, "conversation", "bot", "summary", "ChatGPT", "second", None, "2026-10-09"),
            ChatMessage(3, "conversation", "bot", "chat", "Error", "failed", None, "2026-10-09"),
        ]
        panel.load_conversation()
        self.assertEqual(panel._export_answers, {"answer-1": "first", "answer-2": "second"})
        self.assertIsNone(panel._messages[-1][3])

    def test_chat_render_shows_export_only_for_finished_answers(self):
        panel = object.__new__(ChatPanel)
        panel.history = Mock()
        panel._messages = [("bot", "ChatGPT", "answer", "answer-1"), ("stream", "answering", "partial", None)]
        with patch.object(markdown_html, "show") as show:
            panel._render()
        self.assertEqual(show.call_args.args[1].count("Export…"), 1)
        self.assertIn("rtc-answer-export:answer-1", show.call_args.args[1])

    def test_summary_export_uses_clicked_history_entry(self):
        app = object.__new__(App)
        app.root = Mock()
        app.summary = Mock()
        message = ChatMessage(45, "conversation", "bot", "summary", "ChatGPT", "**summary**", None, "2026-10-09")
        app._summary_messages = {"45": message}
        with patch.object(markdown_html, "show") as show:
            app._show_summary_message(message)
        self.assertIn("rtc-answer-export:45", show.call_args.args[1])
        app._on_summary_link("rtc-answer-export:45")
        with patch.object(export_ui, "choose_export") as choose:
            app.root.after_idle.call_args.args[0]()
        choose.assert_called_once_with(app.root, "**summary**", "chatgpt-summary-45")

    def test_vocab_export_survives_voca_upload_update(self):
        panel = object.__new__(VocabPanel)
        panel._messages, panel._requests = [], {}
        panel._active_request_id = None
        panel._render = panel.set_busy = Mock()
        panel._schedule_render = Mock()
        panel.after_idle = Mock()
        panel._chatgpt_busy = False
        translation = VocabTranslation([NewWord("example", translation="Ví dụ")], "key", "", "")
        panel._on_done(translation)
        message = next(m for m in panel._messages if m[0] == "bot")
        key = message[3]
        panel._on_voca_saved(translation.words, message)
        self.assertEqual(panel._messages[0][3], key)
        panel._on_history_link("rtc-answer-export:" + key)
        with patch.object(export_ui, "choose_export") as choose:
            panel.after_idle.call_args.args[0]()
        self.assertIn("Ví dụ", choose.call_args.args[1])
        self.assertTrue(choose.call_args.kwargs["rendered"])

    def test_cancel_save_dialog_does_not_prepare_or_write_a_file(self):
        with patch.object(export_ui.tk, "Toplevel"), patch.object(export_ui.ttk, "Frame"), \
                patch.object(export_ui.ttk, "Label"), patch.object(export_ui.ttk, "Button") as button, \
                patch.object(export_ui.filedialog, "asksaveasfilename", return_value=""), \
                patch.object(answer_export, "prepare_html") as prepare, patch.object(export_ui, "_start_export") as start:
            export_ui.choose_export(Mock(), "answer", "answer-1")
            button.call_args_list[0].kwargs["command"]()
        prepare.assert_not_called()
        start.assert_not_called()

    def test_format_selection_passes_exact_answer_to_background_export(self):
        with tempfile.TemporaryDirectory() as folder:
            destination = str(Path(folder) / "selected.pdf")
            parent = Mock()
            with patch.object(export_ui.tk, "Toplevel"), patch.object(export_ui.ttk, "Frame"), \
                    patch.object(export_ui.ttk, "Label"), patch.object(export_ui.ttk, "Button") as button, \
                    patch.object(export_ui.filedialog, "asksaveasfilename", return_value=destination) as save, \
                    patch.object(answer_export, "prepare_html", return_value="<p>answer</p>") as prepare, \
                    patch.object(export_ui, "_start_export") as start:
                export_ui.choose_export(parent, "**answer**", "answer-1")
                button.call_args_list[1].kwargs["command"]()
            self.assertEqual(save.call_args.kwargs["defaultextension"], ".pdf")
            prepare.assert_called_once_with("**answer**", rendered=False)
            start.assert_called_once_with(parent, "<p>answer</p>", Path(destination))


if __name__ == "__main__":
    unittest.main()
