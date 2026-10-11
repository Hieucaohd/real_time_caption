"""Offline Word/PDF export of one answer, with shared Markdown/MathJax rendering.

Only embedded PNGs are read; exporting never fetches remote images or contacts
ChatGPT. Conversion uses a temporary sibling file so failures preserve existing
documents. Third-party document libraries are imported only when exporting.
"""

from __future__ import annotations

import base64
import copy
import html
import io
import os
import re
import tempfile
import threading
from dataclasses import dataclass, replace
from pathlib import Path

from . import markdown_html

_LOCK = threading.RLock()
PAGE_WIDTH = 504.0  # Letter, 0.75-inch margins
PAGE_HEIGHT = 684.0


class ExportError(Exception):
    pass


def prepare_html(content: str, *, rendered: bool = False) -> str:
    """Called on the UI thread, where the app's MathJax runtime already lives."""
    result = content if rendered else markdown_html.to_html(content, math_dpi=300)
    if "<code class='tex'>" in result:
        raise ExportError("A formula could not be rendered. Please check its LaTeX before exporting.")
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", result)


@dataclass
class Block:
    tag: str
    node: object
    depth: int = 0
    prefix: str = ""


def _blocks(node, depth=0):
    """Flatten block structure without losing tight lists or mixed inline text."""
    from lxml import etree

    if node.tag in ("p", "pre", "table", "hr", "h1", "h2", "h3", "h4", "h5", "h6"):
        yield Block(node.tag, node, depth)
        return
    if "math-display" in node.get("class", "").split():
        yield Block("math", node, depth)
        return
    if node.tag in ("ul", "ol"):
        start = int(node.get("start", "1"))
        for number, item in enumerate(node, start):
            for index, block in enumerate(_blocks(item, depth + 1)):
                yield replace(block, prefix=(f"{number}. " if node.tag == "ol" else "• ") if index == 0 else block.prefix)
        return
    if node.tag == "blockquote":
        depth += 1
    pending = etree.Element("span")
    pending.text = node.text
    for child in node:
        if child.tag in ("p", "pre", "table", "hr", "ul", "ol", "li", "blockquote", "div", "h1", "h2", "h3", "h4", "h5", "h6"):
            if len(pending) or (pending.text or "").strip():
                yield Block("p", pending, depth)
            pending = etree.Element("span")
            yield from _blocks(child, depth)
            pending.text = child.tail
        else:
            pending.append(copy.deepcopy(child))
    if len(pending) or (pending.text or "").strip():
        yield Block("p", pending, depth)


def _image(node, max_width):
    src = node.get("src", "")
    if not src.startswith("data:image/png;base64,"):
        return None  # No network, local files, or arbitrary URL schemes.
    data = base64.b64decode(src.split(",", 1)[1], validate=True)
    from PIL import Image

    with Image.open(io.BytesIO(data)) as picture:
        w, h = picture.size
    width = float(node.get("width", w)) * 0.75
    height = float(node.get("height", h)) * 0.75
    scale = min(1.0, max_width / max(width, 1), (PAGE_HEIGHT - 24) / max(height, 1))
    return data, width * scale, height * scale


def _rows(node):
    return [list(row) for row in node.xpath("./tr|./thead/tr|./tbody/tr|./tfoot/tr")]


def _docx(root, destination):
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(0.75)
    for name in ("Normal", "Title", *[f"Heading {n}" for n in range(1, 7)]):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    doc.core_properties.title = "ChatGPT response"
    doc.core_properties.author = ""

    def inline(paragraph, node, width, bold=False, italic=False, code=False, strike=False, link=None):
        def text(value):
            if not value:
                return
            run = paragraph.add_run(value)
            run.bold, run.italic, run.font.strike = bold, italic, strike
            if code:
                run.font.name, run.font.size = "Consolas", Pt(10)
            if link:
                hyperlink = OxmlElement("w:hyperlink")
                hyperlink.set(qn("r:id"), paragraph.part.relate_to(link, RT.HYPERLINK, is_external=True))
                run.font.color.rgb = RGBColor(0x1A, 0x5F, 0xB4)
                run.font.underline = True
                hyperlink.append(run._r)
                paragraph._p.append(hyperlink)

        text(node.text)
        for child in node:
            if child.tag == "img":
                picture = _image(child, width)
                if picture:
                    data, w, h = picture
                    run = paragraph.add_run()
                    shape = run.add_picture(io.BytesIO(data), width=Pt(w), height=Pt(h))
                    shape._inline.docPr.set("descr", child.get("alt", "Formula"))
                    if "math-inline" in child.get("class", ""):
                        match = re.search(r"vertical-align:([+-]?[\d.]+)px", child.get("style", ""))
                        position = OxmlElement("w:position")
                        position.set(qn("w:val"), str(round(float(match[1]) * 1.5) if match else -4))
                        run._r.get_or_add_rPr().append(position)
                else:
                    text(child.get("alt", "[Image]"))
            elif child.tag == "br":
                paragraph.add_run().add_break()
            else:
                inline(paragraph, child, width, bold or child.tag in ("strong", "b"),
                       italic or child.tag in ("em", "i"), code or child.tag == "code",
                       strike or child.tag == "s", child.get("href") if child.tag == "a" else link)
            text(child.tail)

    def write(container, blocks, width):
        for block in blocks:
            if block.tag == "table":
                rows = _rows(block.node)
                if not rows:
                    continue
                count = max(map(len, rows))
                table = container.add_table(rows=len(rows), cols=count)
                table.style = "Table Grid"
                table.autofit = False
                for column in table.columns:
                    column.width = Pt(width / count)
                for ri, row in enumerate(rows):
                    for ci, cell_node in enumerate(row):
                        cell = table.cell(ri, ci)
                        cell.width = Pt(width / count)
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        # Empty initial cell paragraph is removed after real content is added.
                        initial = cell.paragraphs[0]._p
                        write(cell, _blocks(cell_node), width / count - 12)
                        if len(cell.paragraphs) > 1:
                            cell._tc.remove(initial)
                        if cell_node.tag == "th":
                            shade = OxmlElement("w:shd")
                            shade.set(qn("w:fill"), "F3F3F3")
                            cell._tc.get_or_add_tcPr().append(shade)
                            for p in cell.paragraphs:
                                for r in p.runs:
                                    r.bold = True
                    if ri == 0 and any(c.tag == "th" for c in row):
                        table.rows[ri]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                # Explicit light gray borders rather than Word's default black grid.
                borders = OxmlElement("w:tblBorders")
                for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                    border = OxmlElement(f"w:{edge}")
                    for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                        border.set(qn(f"w:{key}"), value)
                    borders.append(border)
                table._tbl.tblPr.append(borders)
                continue
            if block.tag == "hr":
                container.add_paragraph()
                continue
            p = container.add_paragraph(style=f"Heading {block.tag[1]}" if block.tag.startswith("h") else None)
            p.paragraph_format.left_indent = Pt(block.depth * 18)
            available = max(24, width - block.depth * 18)
            if block.prefix:
                p.add_run(block.prefix)
            if block.tag == "pre":
                r = p.add_run("".join(block.node.itertext()).rstrip("\n").expandtabs(4))
                r.font.name, r.font.size = "Consolas", Pt(10)
            else:
                inline(p, block.node, available)
            if block.tag == "math":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(8)

    write(doc, _blocks(root), PAGE_WIDTH)
    doc.save(destination)


def _pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "CaptionBody" in pdfmetrics.getRegisteredFontNames():
        return
    candidates = [
        (Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts", ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf", "consola.ttf")),
        (Path("/usr/share/fonts/truetype/dejavu"), ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSans-BoldOblique.ttf", "DejaVuSansMono.ttf")),
    ]
    for folder, files in candidates:
        if all((folder / f).exists() for f in files):
            for name, filename in zip(("CaptionBody", "CaptionBold", "CaptionItalic", "CaptionBoldItalic", "CaptionCode"), files):
                pdfmetrics.registerFont(TTFont(name, str(folder / filename)))
            pdfmetrics.registerFontFamily("CaptionBody", normal="CaptionBody", bold="CaptionBold", italic="CaptionItalic", boldItalic="CaptionBoldItalic")
            return
    raise ExportError("No Unicode font found for PDF. Install Arial or DejaVu Sans and try again.")


def _pdf(root, destination):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Image, LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    _pdf_fonts()
    body = ParagraphStyle("Body", fontName="CaptionBody", fontSize=11, leading=16, spaceAfter=6, autoLeading="max")

    def inline(node, width):
        parts = [html.escape(node.text or "")]
        tags = {"strong": "b", "b": "b", "em": "i", "i": "i", "s": "strike"}
        for child in node:
            if child.tag == "img":
                image = _image(child, width)
                if image:
                    _, w, h = image
                    parts.append(f'<img src="{child.get("src")}" width="{w}" height="{h}" valign="middle"/>')
                else:
                    parts.append(html.escape(child.get("alt", "[Image]")))
            elif child.tag == "br":
                parts.append("<br/>")
            elif child.tag in tags:
                tag = tags[child.tag]
                parts.append(f"<{tag}>{inline(child, width)}</{tag}>")
            elif child.tag == "code":
                parts.append(f'<font name="CaptionCode">{inline(child, width)}</font>')
            elif child.tag == "a" and child.get("href", "").startswith(("https://", "http://", "mailto:")):
                parts.append(f'<link href="{html.escape(child.get("href"), quote=True)}" color="#1a5fb4">{inline(child, width)}</link>')
            else:
                parts.append(inline(child, width))
            parts.append(html.escape(child.tail or ""))
        return "".join(parts)

    def write(blocks, width):
        result = []
        for block in blocks:
            if block.tag == "math":
                # A tall display equation must be a standalone flowable. Putting
                # it in a middle-aligned inline <img> can overlap the prior line.
                for node in block.node.xpath(".//img"):
                    image = _image(node, width)
                    if image:
                        data, w, h = image
                        picture = Image(io.BytesIO(data), width=w, height=h, hAlign="CENTER")
                        picture.spaceBefore = picture.spaceAfter = 8
                        result.append(picture)
                continue
            if block.tag == "table":
                rows = _rows(block.node)
                if not rows:
                    continue
                count = max(map(len, rows))
                data = []
                for row in rows:
                    data.append([write(_blocks(cell), width / count - 16) or [Paragraph("", body)] for cell in row] + [""] * (count - len(row)))
                header = any(c.tag == "th" for c in rows[0])
                table = LongTable(data, colWidths=[width / count] * count, repeatRows=int(header), splitInRow=1, hAlign="LEFT", spaceAfter=8)
                styles = [("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                          ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                          ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                          ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]
                if header:
                    styles.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3F3F3")))
                table.setStyle(TableStyle(styles))
                result.append(table)
                continue
            if block.tag == "hr":
                result.append(Spacer(1, 10))
                continue
            style = ParagraphStyle("Block", parent=body, leftIndent=block.depth * 18)
            available = max(24, width - block.depth * 18)
            if block.tag.startswith("h"):
                style.fontName = "CaptionBold"
                style.fontSize = max(12, 22 - int(block.tag[1]) * 2)
                style.leading = style.fontSize * 1.3
                style.spaceBefore, style.keepWithNext = 10, True
            if block.tag == "pre":
                style.fontName, style.fontSize, style.leading = "CaptionCode", 10, 14
                value = html.escape("".join(block.node.itertext()).rstrip("\n").expandtabs(4))
                value = value.replace(" ", "&#160;").replace("\n", "<br/>")
            else:
                value = inline(block.node, available)
            result.append(Paragraph(html.escape(block.prefix) + value, style))
        return result

    doc = SimpleDocTemplate(str(destination), pagesize=letter, topMargin=54, bottomMargin=54, leftMargin=54, rightMargin=54, title="ChatGPT response", author="")
    # SimpleDocTemplate's frame also has 6-point inner padding on each side.
    doc.build(write(_blocks(root), PAGE_WIDTH - 12))


def export_html(content: str, destination: Path) -> None:
    """Export prepared HTML to .docx/.pdf; safe to call in a background worker."""
    destination = Path(destination)
    extension = destination.suffix.lower()
    if extension not in (".docx", ".pdf"):
        raise ExportError("Choose a Word (.docx) or PDF (.pdf) filename.")
    try:
        from lxml.html import fragment_fromstring
    except ImportError as exc:
        raise ExportError("Export libraries are missing. Run setup.bat to install them.") from exc
    root = fragment_fromstring(content or "<p></p>", create_parent="div")
    temp_path = None
    try:
        with _LOCK:
            fd, filename = tempfile.mkstemp(prefix=".caption-export-", suffix=extension, dir=destination.parent)
            os.close(fd)
            temp_path = Path(filename)
            (_docx if extension == ".docx" else _pdf)(root, temp_path)
            os.replace(temp_path, destination)
    except ImportError as exc:
        raise ExportError("Export libraries are missing. Run setup.bat to install python-docx and reportlab.") from exc
    finally:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)
